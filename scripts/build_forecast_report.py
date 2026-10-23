"""Build a concise diagnostic report after independent artifact audit passes."""
from pathlib import Path
from datetime import datetime, timezone, timedelta
import json
import statistics as stats
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / 'results/gas-marl-20261023-forecast-pilot-01'
OUT = ROOT / 'reports'
proof = json.loads((ROOT / 'research/evidence/forecast_pilot_01/proof.json').read_text(encoding='utf-8'))
assert proof['status'] == 'verified'
summary = json.loads((RUN / 'summary.json').read_text(encoding='utf-8'))
rows = json.loads((RUN / 'rows.json').read_text(encoding='utf-8'))
assert summary['status'] == 'completed' and summary['h1']['protocol_valid'] and summary['h1']['supported']
assert len(rows) == summary['episodes_completed'] == summary['episodes_started'] == 18
by = {(r['seed'], r['start'], r['condition']): r for r in rows}
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Cm(1.6), Cm(1.6)
section.left_margin, section.right_margin = Cm(2), Cm(2)
for name, size in [('Normal', 11), ('Title', 21), ('Subtitle', 10), ('Heading 1', 14)]:
    style = doc.styles[name]
    style.font.name, style.font.size = 'Calibri', Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)
for item in list(doc.styles.element.iter(qn('w:pBdr'))):
    item.getparent().remove(item)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('GAS-MARL 冻结模型先导  ·  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
md = ['# GAS MARL 预测偏差先导诊断报告', '']


def p(text):
    md.extend([text, ''])
    return doc.add_paragraph(text)


def h(text, page=False):
    item = doc.add_heading(text, 1)
    item.paragraph_format.page_break_before = page
    md.extend(['## ' + text, ''])


def table(headers, values, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, headers): cell.text = text
    for values_ in values:
        for cell, text in zip(t.add_row().cells, values_): cell.text = str(text)
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                element = OxmlElement('w:' + edge)
                for key, value_ in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: element.set(qn('w:' + key), value_)
                borders.append(element)
            pr.append(borders); margins = OxmlElement('w:tcMar')
            for edge, number in [('top', '90'), ('bottom', '90'), ('left', '90'), ('right', '90')]:
                element = OxmlElement('w:' + edge); element.set(qn('w:w'), number); element.set(qn('w:type'), 'dxa'); margins.append(element)
            pr.append(margins); shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCE6F1' if ri == 0 else ('F7F9FB' if ri % 2 == 0 else 'FFFFFF')); pr.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0); paragraph.paragraph_format.line_spacing = 1.1
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs: run.font.size = Pt(9.5); run.bold = ri == 0
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    cell_md = lambda value: str(value).replace('|', '&#124;').replace('\n', '<br>')
    md.extend(['| ' + ' | '.join(map(cell_md, headers)) + ' |', '|' + '---|' * len(headers)])
    md.extend('| ' + ' | '.join(map(cell_md, row)) + ' |' for row in values); md.append('')


def mean(seed, condition, key):
    return stats.mean(by[seed, start, condition][key] for start in (6567, 919))


doc.add_paragraph('GAS MARL 预测偏差先导诊断报告', 'Title')
doc.add_paragraph('三个冻结模型  两个固定窗口  未来预测高估20%  2026年10月23日', 'Subtitle')
ended = datetime.fromisoformat(summary['ended_utc']).astimezone(timezone(timedelta(hours=8)))
p(f'先导于北京时间{ended:%H时%M分}完成，18次完整评测全部结束。六组原始路径与零偏差路径在动作、事件、排程、逐作业记录、物理指标和随机状态上完全一致；随后才运行六次偏差条件。没有训练新模型，也没有追加幅度、种子或窗口。')
h('一 结果支持什么')
p('在本次人为压力条件下，三个种子的两窗口平均真实ReUtil均下降，支持预先定义的H1。但六个窗口配对中有一个ReUtil提高，BSLD和综合奖励也呈现取舍；这个结果只证明有限条件下的预测敏感性，不能写成所有场景性能退化、稳健算法有效或创新已成立。')
data = []
for seed in (0, 1, 2):
    data.append([seed, f"{100*mean(seed,'zero','renewable_energy_utilization'):.2f}", f"{100*mean(seed,'plus20','renewable_energy_utilization'):.2f}", f"{100*(mean(seed,'plus20','renewable_energy_utilization')-mean(seed,'zero','renewable_energy_utilization')):+.3f}", f"{mean(seed,'zero','average_bounded_slowdown'):.2f} → {mean(seed,'plus20','average_bounded_slowdown'):.2f}", f"{mean(seed,'zero','reward'):.4f} → {mean(seed,'plus20','reward'):.4f}"])
table(['种子', '基线绿电\n%', '高估后绿电\n%', '绿电变化\n百分点', '平均BSLD\n基线 → 高估', '综合奖励\n基线 → 高估'], data, [1.3, 2.4, 2.5, 2.5, 4, 4.1])
p('例如种子1在窗口6567的绿电利用率下降约3.36个百分点，BSLD却从35.25降至12.50，综合奖励增加约0.01194。不同指标的方向不能互相替代；报告和源数据保留全部正负变化。')
h('二 固定条件与隔离闸门', page=True)
p('模型是已有第300轮GAS-MARL权重，保持author语义、Lublin-256和Green-Backfilling。按原评测清单先选6567，再跳过重叠窗口7146，选919；保留原iteration 1和3进行paired播种。它们是复用窗口，不是新测试集。')
p('仅把当前小时之后的供给预测乘1.20，当前已观测小时保持准确；24槽状态、作业绿色特征和回填估计均经过预测接口。真实供给积分、作业和功率保持固定。零偏差保留原窗口重复规则，不等同于一个经过现实误差标定的预报器。')
p('执行使用共享研究锁、新输出目录、550秒软超时和600秒进程树硬超时。源代码、输入与最终权重在运行前后校验；预测接口和执行器通过110项整合测试后才运行。完整事件轨迹保留本地，小型证据提供哈希与独立复核结果。')
h('三 六组窗口的代价与取舍')
data = []
for seed in (0, 1, 2):
    for start in (6567, 919):
        a, b = by[seed,start,'zero'], by[seed,start,'plus20']
        data.append([f'{seed} / {start}', f"{a['average_bounded_slowdown']:.2f} → {b['average_bounded_slowdown']:.2f}", f"{100*(b['renewable_energy_utilization']-a['renewable_energy_utilization']):+.3f}", f"{b['reward']-a['reward']:+.5f}", f"{(b['brown_energy_j']-a['brown_energy_j'])/1e6:+.2f}", f"{(b['wait_quantiles']['p95']-a['wait_quantiles']['p95'])/3600:+.2f}"])
table(['种子 / 起点', 'BSLD\n基线 → 高估', '绿电变化\n百分点', '奖励变化', '棕电变化\nMJ', '等待P95\n变化 小时'], data, [2.4, 3.4, 2.4, 2.7, 2.9, 3])
p('正的等待或棕电变化表示代价增加。分位数为每个episode内1024个作业的线性插值分位数；没有把作业或窗口当作独立训练重复，也没有进行显著性检验。逐条件绝对值、P50/P95/P99及最大值、三指标和总绿棕能耗均保存在源数据。')
h('四 资源开销与测量边界', page=True)
rss = max(r['process_memory']['peak_rss_bytes'] for r in rows)
gpu = max(r['gpu']['peak_allocated_bytes'] for r in rows)
p(f"评测矩阵耗时{summary['total_wall_seconds']:.2f}秒，进程CPU时间{summary['total_process_cpu_seconds']:.2f}秒。观测到的进程峰值工作集约{rss/1024**3:.2f}GiB，CUDA分配器峰值约{gpu/1024**2:.2f}MiB，日志统计约{summary['log_bytes_before_final_summary']/1e6:.2f}MB；前两者不是整机RAM或显卡全部占用。GPU利用率没有采样，不能由分配显存推算。")
p('能耗采用作者的首作业启动至最终实际完成的PowerStruc积分范围，并单独记录首到达至首启动的空闲间隔与静态能耗；没有悄悄修正原ReUtil分母。makespan按首到达至实际最终完成计算。这里的真实能耗指仿真真实供给下的集群积分，不是本机插座耗电。')
p('源数据同时保存动作、请求与实际延迟、到达backfill_check的候选及接受次数、实际回填数和预测上限裁切。裁切比例的观测时点会随排程改变，不能把其相关变化当作延迟或回填单一通道的因果消融。原始全数组前后哈希字符串未单独保存；独立审计使用冻结源数据重建物理积分，并区分它与执行时的数组不变断言。')
h('五 下一阶段需要的证据')
p('目前足以继续调查预测误差及多目标取舍，但尚不能据此确定SCI新方法。GAS-MARL-ADC和SCS-RC直接近邻的合法全文仍未取得，观测权限、动作、约束、误差注入与评测划分尚有未知。下一步优先补齐这两篇的合法作者稿或授权全文，再判断候选是否重复及是否值得提出最小改造；本次不追加长期训练。')
p('原始来源：https://doi.org/10.2139/ssrn.6530554；https://doi.org/10.1016/j.apenergy.2026.128137。复核入口：scripts/analyze_forecast_pilot.py；证据：research/evidence/forecast_pilot_01。')
OUT.mkdir(exist_ok=True)
doc.save(OUT/'预测偏差先导报告_GAS_MARL.docx')
(OUT/'预测偏差先导报告_GAS_MARL.md').write_text('\n'.join(md),encoding='utf-8')
print('reports/预测偏差先导报告_GAS_MARL.docx')
