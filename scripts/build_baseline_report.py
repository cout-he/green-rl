"""Build the audited first baseline report using the bundled python-docx runtime."""
from pathlib import Path
import csv
from datetime import datetime, timezone, timedelta
import hashlib
import json
import shutil
import statistics
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / 'results/gas-marl-20261021-baseline-02'
PAIR = ROOT / 'results/gas-marl-20261022-paired-01'
OUT = ROOT / 'reports'
EVIDENCE = ROOT / 'research/evidence/baseline02'
summary = json.loads((RUN / 'analysis/summary.json').read_text(encoding='utf-8'))
verification = json.loads((RUN / 'analysis/verification.json').read_text(encoding='utf-8'))
state = json.loads((RUN / 'state.json').read_text(encoding='utf-8'))
pair_state = json.loads((PAIR / 'state.json').read_text(encoding='utf-8'))
assert verification['status'] == 'verified' and pair_state['status'] == 'completed' and pair_state['control_equality']
author = {r['algorithm']: r for r in summary['evaluations'] if r['semantics'] == 'author'}
paired = {}
for mode in ('author', 'paper'):
    rows = json.loads((PAIR / mode / 'rows.json').read_text())
    paired[mode] = {algorithm: {key: statistics.mean(r[key] for r in rows if r['algorithm'] == algorithm)
                              for key in ('average_bounded_slowdown', 'renewable_energy_utilization')}
                    for algorithm in ('FCFS', 'F2', 'PPO', 'MARL')}
legacy = {mode: json.loads((RUN / 'baseline/lublin_256/seed_0/evaluation' / mode / 'rows.json').read_text())
          for mode in ('author', 'paper')}
legacy_ppo_equal = sum(a == b for a, b in zip([r for r in legacy['author'] if r['algorithm'] == 'PPO'],
                                             [r for r in legacy['paper'] if r['algorithm'] == 'PPO']))
assert legacy_ppo_equal == 1
bsld_gain = 100 * (author['MARL']['bsld_mean'] / author['PPO']['bsld_mean'] - 1)
reutil_gain = 100 * (author['MARL']['reutil_mean'] - author['PPO']['reutil_mean'])
pa, pp = paired['author']['MARL'], paired['paper']['MARL']
paired_bsld_pct = 100 * (pp['average_bounded_slowdown'] / pa['average_bounded_slowdown'] - 1)
paired_reutil_pp = 100 * (pp['renewable_energy_utilization'] - pa['renewable_energy_utilization'])
completed = datetime.fromisoformat(state['ended_at']).astimezone(timezone(timedelta(hours=8)))
hours = (datetime.fromisoformat(state['ended_at']) - datetime.fromisoformat(state['started_at'])).total_seconds() / 3600

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.7)
section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
for name, size in [('Normal', 11), ('Title', 21), ('Subtitle', 10), ('Heading 1', 14), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.paragraph_format.line_spacing = 1.2
    style.paragraph_format.space_after = Pt(7)
for name in ('Heading 1', 'Heading 2'):
    doc.styles[name].paragraph_format.space_before = Pt(8)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('GAS-MARL | 首轮基线验证  ·  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
md = ['# GAS-MARL 首轮基线验证报告', '']


def p(text):
    paragraph = doc.add_paragraph(text)
    md.extend([text, ''])
    return paragraph


def h(text, page=False):
    paragraph = doc.add_heading(text, 1)
    if page:
        paragraph.paragraph_format.page_break_before = True
    md.extend(['## ' + text, ''])


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    for row_index, row in enumerate(t.rows):
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement('w:' + edge)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    item.set(qn('w:' + key), value)
                borders.append(item)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge, value in [('top', '80'), ('bottom', '80'), ('left', '100'), ('right', '100')]:
                item = OxmlElement('w:' + edge)
                item.set(qn('w:w'), value)
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            pr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DCE6F1' if row_index == 0 else ('F7F9FB' if row_index % 2 == 0 else 'FFFFFF'))
            pr.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    run.bold = row_index == 0
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    md.extend(['| ' + ' | '.join(headers) + ' |', '|' + '---|' * len(headers)])
    md.extend('| ' + ' | '.join(map(str, row)) + ' |' for row in rows)
    md.append('')


def figure(stem, caption):
    image = OUT / 'figures' / (stem + '.png')
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image), width=Inches(6.5))
    paragraph.paragraph_format.keep_with_next = True
    cap = p(caption)
    for run in cap.runs:
        run.font.size = Pt(9)
    md.extend([f'![{caption}](figures/{stem}.png)', ''])


doc.add_paragraph('GAS-MARL 首轮基线验证报告', 'Title')
doc.add_paragraph('Lublin-256 · Green-Backfilling · 训练种子0 · Python/PyTorch', 'Subtitle')
p(f'完整流水线完成于北京时间{completed:%Y年%m月%d日 %H:%M}，用时{hours:.2f}小时。当前交付为一组独立训练种子的完整基线及固定权重语义对照。')
h('1. 结论与论文对照')
p('作者模式的主要性能排序得到支持：本次GAS-MARL在平均BSLD和绿电利用率上均优于PPO。不过，单个训练种子、软件版本偏离和随机评测不足以证明论文数值已复现，也不构成新的算法贡献。')
rows = []
for algorithm in ('FCFS', 'F2', 'PPO', 'MARL'):
    r = author[algorithm]
    anchor = r.get('published_anchor')
    rows.append(['GAS-MARL' if algorithm == 'MARL' else algorithm, f"{r['bsld_mean']:.4f}",
                 f"{100*r['reutil_mean']:.2f}", f"{anchor['bsld']:.3f}" if anchor else '—',
                 f"{100*anchor['reutil']:.2f}" if anchor else '—'])
table(['算法', '本次BSLD↓', '本次绿电/%↑', '论文BSLD', '论文绿电/%'], rows, [3.1, 3.1, 3.2, 3.0, 3.2])
p(f'相对本次PPO，GAS-MARL的BSLD低{abs(bsld_gain):.2f}%，绿电利用率高{reutil_gain:.2f}个百分点。此比较描述本次相同十段输入上的结果，不应解释为对论文方法的新改进。')
p('相对论文Table 3/4，GAS-MARL的BSLD偏差为−37.49%，ReUtil偏差为−1.20个百分点；PPO分别为+15.82%与−0.89个百分点。BSLD降低与绿电利用率降低同时存在，不能只挑选有利指标判定复现成功。')
p('指标方向：BSLD越低越好，ReUtil越高越好。ReUtil的分母是集群总耗能；表中百分数不是“可用绿电消纳率”。')
h('2. 训练轨迹与逐序列结果', page=True)
figure('gas_marl_baseline_training', '图1｜两算法全部300轮训练记录。浅线保留每轮均值，实线为后向10轮移动均值；首9轮使用已有轮数。只有一个独立训练种子，不绘制种子间置信区间。')
figure('gas_marl_baseline_author_evaluation', '图2｜作者模式的十段1024作业评测。保留各段原始点，均值和样本SD仅描述片段差异；空心标记为论文锚点。十个作业窗口中有11对重叠，不把这些点视为十次独立训练。')
h('3. 随机采样混杂与隔离对照', page=True)
p('原评测程序只在整次评测开始播种一次，随后不同算法和不同片段共用随机流。截止语义改变MARL的轨迹长度后，会改变后续PPO所用的随机数。原始两模式的PPO第一段相同、后九段不同；PPO的延迟动作固定为零，因此其差异不能归因于截止修正。')
p('为隔离该混杂，保留全部原始结果和最终第300轮权重，新增paired模式：每个算法、每段起点分别重置Python、NumPy、Torch及环境随机状态，两种语义使用同一派生种子。没有重新训练模型，也没有覆盖原评测。')
pair_rows = []
for algorithm in ('PPO', 'MARL'):
    a, b = paired['author'][algorithm], paired['paper'][algorithm]
    pair_rows.append(['GAS-MARL' if algorithm == 'MARL' else algorithm,
                      f"{a['average_bounded_slowdown']:.4f} → {b['average_bounded_slowdown']:.4f}",
                      f"{100*a['renewable_energy_utilization']:.2f} → {100*b['renewable_energy_utilization']:.2f}"])
table(['固定模型', 'BSLD：author → paper', '绿电/%：author → paper'], pair_rows, [3.2, 6.2, 6.2])
p(f'PPO、FCFS、F2在两模式下逐行完全相同，零效应对照通过。GAS-MARL的BSLD变化为{paired_bsld_pct:+.2f}%，绿电利用率变化为{paired_reutil_pp:+.2f}个百分点；逐段BSLD有改善、恶化和不变，不能声称所有序列均受益。')
p('隔离对照的方向与原始共享随机流比较不同，说明修复评测设计会影响研究结论。新的author结果与原始author结果也可不同，因为采样种子方案发生改变；两套数字不能混成一个表或挑选最优值。')
p('限制：这仍是同一个作者模式训练的模型、一个评测基础种子。它验证的是Type 2截止事件的语义敏感性，不是独立训练的新方法；需要更多训练和评测种子才能评估稳定性。未进行显著性检验。')
h('4. 证据、版本与后续安排', page=True)
table(['核验项目', '已核实结果'], [
    ['训练完整性', '两算法各300行CSV、300条连续epoch日志；各行数值一致'],
    ['权重与恢复状态', '最终checkpoint均为第300轮；网络与独立权重逐tensor一致，参数有限，双Adam与全部RNG齐全'],
    ['运行与来源', '8个步骤退出0、stderr均为空；23个源码快照和6个数据文件SHA256通过'],
    ['评测对齐', '每模式4算法×10段×1024作业；逐序列均值、CSV与奖励公式一致'],
], [3.4, 12.2])
p('训练源版本：0b6541e；作者参考版本：0d89df5。原作者导入保留于upstream-import。配对评测使用另存的比较脚本和原始权重SHA256清单，训练权重未被改写。原01任务的127轮中断数据与本次完整02任务分开保留。')
p('环境为Python 3.11.15、PyTorch 2.7.1+cu128、NumPy 1.24.4、RTX 5060。相对作者依赖，PyTorch由2.1.1升级以支持显卡，NumPy冲突固定版本统一为1.24.4，OpenCV固定4.11.0.86；这些偏离已进入运行环境证据。绘图依赖在训练完成后安装，未冒充为训练时环境。')
p('工程状态：连续恢复时已完成算法的检查点传递缺陷已补失败测试并修复，不覆盖新检查点，复制失败会明确标记运行失败。该改动属于可靠性修复，不算创新。')
p('下一阶段优先验证跨随机种子的稳定性，再决定是否进入绿电预测误差的创新试验。现有近邻文献已覆盖预测、注意力、双critic和部分风险调度，尚缺关键全文，不能仅凭概念名称声称新颖。新增长期实验需记录独立种子与一致评测协议。')
p('自动定时监测按用户要求保持关闭。代码、小型证据、逐序列数据和可编辑报告同步研究分支；训练权重、虚拟环境及用户论文PDF留在本地。')
p('证据索引：研究目录下的baseline02证据包含原始训练CSV、两套原始逐序列评测、配对评测、核验记录及源文件/数据哈希；图表附源数据CSV与绘图脚本。')

OUT.mkdir(exist_ok=True)
doc.save(OUT / '基线验证报告_GAS_MARL.docx')
(OUT / '基线验证报告_GAS_MARL.md').write_text('\n'.join(md), encoding='utf-8')
EVIDENCE.mkdir(parents=True, exist_ok=True)
for relative in ('source_manifest.json', 'datasets.json', 'environment.json', 'analysis/verification.json'):
    shutil.copyfile(RUN / relative, EVIDENCE / Path(relative).name)
for folder in ('MARL', 'MaskablePPO'):
    shutil.copyfile(RUN / 'baseline/lublin_256/seed_0/models/lublin_256' / folder / f'{folder}_lublin_256.csv', EVIDENCE / f'{folder}_training.csv')
for mode in ('author', 'paper'):
    for name in ('rows.json', 'metadata.json', 'result.csv'):
        shutil.copyfile(RUN / 'baseline/lublin_256/seed_0/evaluation' / mode / name, EVIDENCE / f'legacy_{mode}_{name}')
        shutil.copyfile(PAIR / mode / name, EVIDENCE / f'paired_{mode}_{name}')
record = {'status': 'completed_first_seed_baseline', 'training_seed_count': 1, 'training_seed': 0,
          'run_id': RUN.name, 'source_commit': state['git_commit'], 'completed_at_beijing': completed.isoformat(),
          'elapsed_hours': hours, 'author': author, 'paired_means': paired,
          'paired_controls_exactly_equal': True, 'legacy_ppo_equal_windows': legacy_ppo_equal,
          'paired_marl_bsld_change_percent': paired_bsld_pct, 'paired_marl_reutil_change_pp': paired_reutil_pp,
          'model_sha256': pair_state['model_sha256'],
          'paired_compare_sha256': hashlib.sha256((PAIR / 'compare.py').read_bytes()).hexdigest(),
          'limits': verification['limits'] + ['Original paper-author comparison is confounded by shared RNG consumption.', 'Matched evaluation uses one base seed.']}
(EVIDENCE / 'baseline_summary.json').write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding='utf-8')
print(OUT / '基线验证报告_GAS_MARL.docx')
