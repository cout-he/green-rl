"""Build the Chinese template A report with the bundled python-docx runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports"
OUT.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.1), Cm(2.1)

for name, size in (("Normal", 10.5), ("Title", 22), ("Subtitle", 10),
                   ("Heading 1", 15), ("Heading 2", 12)):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(6)
doc.styles["Title"].paragraph_format.space_after = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(7)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
for borders in list(doc.styles.element.iter(qn("w:pBdr"))):
    borders.getparent().remove(borders)

markdown = []
new_page = False

def p(text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    markdown.append(text + "\n")
    return paragraph

def h(text, level=1):
    global new_page
    heading = doc.add_heading(text, level)
    if new_page:
        heading.paragraph_format.page_break_before = True
        new_page = False
    markdown.append("#" * (level + 1) + " " + text + "\n")

def equation(text, label):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    content = OxmlElement("m:t")
    content.text = text
    run.append(content)
    math.append(run)
    paragraph._p.append(math)
    paragraph.add_run("    " + label).font.size = Pt(9)
    markdown.append("$$" + text + "$$  " + label + "\n")

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, header in enumerate(headers):
        t.rows[0].cells[i].text = header
    for row in rows:
        for i, value in enumerate(row):
            if i == 0:
                cells = t.add_row().cells
            cells[i].text = value
    for row_index, row in enumerate(t.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement("w:" + edge)
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), "D9D9D9")
                borders.append(el)
            tc_pr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge, value in (("top", "65"), ("bottom", "65"), ("left", "100"), ("right", "100")):
                el = OxmlElement("w:" + edge)
                el.set(qn("w:w"), value)
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            tc_pr.append(margins)
            fill = "DCE6F1" if row_index == 0 else ("F7F9FB" if row_index % 2 == 0 else "FFFFFF")
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), fill)
            tc_pr.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                if col_index > 0 and len(headers) == 5:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = row_index == 0
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.add_run().font.size = Pt(2)
    markdown.append("| " + " | ".join(headers) + " |")
    markdown.append("|" + "---|" * len(headers))
    markdown.extend("| " + " | ".join(row) + " |" for row in rows)
    markdown.append("")

def page():
    global new_page
    new_page = True
    markdown.append("\n---\n")

doc.add_paragraph("GAS MARL 模型拆解与研究方向报告", "Title")
markdown.append("# GAS MARL 模型拆解与研究方向报告\n")
p("文献  Rui Chen 等  Future Generation Computer Systems 167 107760 2025", bold_lead="文献")
p("领域  高性能计算作业调度  多动作深度强化学习  可再生能源利用", bold_lead="领域")
p("本研究建议先复现作者的双目标调度基线，再优先评估可再生能源预测误差对调度的影响。源码具备模型、仿真和数据主体，可以采用 Python 与 PyTorch 工作流；CUDA短跑链条已通过，尚未完整复现论文数值，工程测试不能替代研究结论。")

h("1 模型讲解")
p("把集群想成一间有固定数量座位的计算机房。作业陆续到达，说明需要多少处理器、预计运行多久及功率。太阳能与风能会随小时变化，排得早不一定更绿色，等得久又会损害用户体验。GAS-MARL 每次先选择一个作业，再决定立即尝试、等待若干运行作业完成，或延迟固定时长。这里 MARL 指 Multi-Action，并非多智能体。")
p("Green-Backfilling 在选中作业等待时，让合适的小作业填补空闲资源。它优先考虑请求时间、处理器数和功率乘积较小的作业，并限制新增棕色能源消耗，从而尽量减少资源闲置和额外排放。")
h("核心假设与输入", 2)
p("集群同构，处理器为最小资源单位；每台机器有8个处理器、50 W静态功耗，运行作业带来附加功耗。作业不可抢占，功率在运行期间恒定；棕色供给充足，绿色优先。模型可获得未来24个一小时槽的能源预测，超过预测窗口时重复对应槽。")
p("输入包括三种作业轨迹 Lublin-256、Cirne、Jann，作业提交时间、请求资源、请求时间和真实运行时间，以及每处理器功率、风速和辐照数据。状态使用256个等待作业、64个运行作业和24个绿色供给槽；训练默认300轮，每轮100条、每条256个作业。")
h("目标与关键公式", 2)
equation("P_C(s) = P_idle + Σⱼ∈R(s) pow_j", "原文式1")
equation("ReUtil = ∫ min(P_GT(s), P_C(s)) ds / ∫ P_C(s) ds", "原文式2至4")
equation("AvgBSLD = (1/N) Σᵢ max((wᵢ + eᵢ) / max(10, eᵢ), 1)", "原文式5")
equation("r_done = ReUtil − η × AvgBSLD     r_other = 0", "原文式9")
p("P_idle 为全机房静态功率，pow_j 为作业附加功率；P_GT 为绿色供给。w、e 分别为等待和真实运行时间，N 为作业数。ReUtil 越大越好，AvgBSLD 越小越好；η 调节等待代价。输出是作业选择、延迟决策、调度日志及两项指标，不是直接输出最优解证明。")

page()
h("2 拆解清单与复现锚点")
table(["类别", "内容与原文位置", "源码位置及待补项"], [
    ["输入", "作业与功率数据  第5.1节 式16至18", "data/ 与 job.py\n保留原始数据及预处理口径"],
    ["状态与动作", "三类状态 双动作  第4.2节 图2", "HPCSimPickJobs.py 与 MARL.py\n核对掩码和维度"],
    ["学习", "PPO 稀疏奖励  式9至14 算法1", "MARL.py 与 MaskablePPO.py\n需实际训练及种子复核"],
    ["回填", "优先级与50000 J门槛  式15 算法2", "HPCSimPickJobs.py 与 cluster.py\nType2截止语义单列验证"],
    ["指标", "能源积分与slowdown  式1至5", "PowerStruc.py 与 greenPower.py\n使用实际运行时积分与评分"],
    ["输出", "均值 训练曲线 策略比较  图3至7 表3至4", "compare.py\n需训练权重 逐序列结果与统计汇总"],
], [2.1, 6.6, 8.1])
h("优先核对的原文数值", 2)
p("下表仅摘录论文第10页表3和表4的 Green-Backfilling 条件，作为复现目标，均不是本项目实验结果。每种负载评估10条长度1024的序列，PPO与GAS-MARL应使用对应回填策略训练和评估。")
table(["负载", "PPO\nAvgBSLD", "GAS\nAvgBSLD", "PPO\nReUtil", "GAS\nReUtil"], [
    ["Lublin-256", "56.358", "31.738", "0.6767", "0.8154"],
    ["Cirne", "22.357", "17.895", "0.5311", "0.7545"],
    ["Jann", "15.374", "12.534", "0.5454", "0.7424"],
], [4.4, 3.1, 3.1, 3.1, 3.1])
p("解释数值时应保留双目标取舍：表3中，EASY条件下 Lublin-256 的GAS-MARL slowdown为63.627，高于PPO的55.262；Green条件下 Jann 的F2为11.332，低于GAS-MARL的12.534。不能把论文理解为所有工况、所有指标均由GAS-MARL最优。")
h("作者模式与修正模式", 2)
p("原始导入以 upstream-import 标签追溯，工作副本关联 cout-he/green-rl。author保留作者语义；paper仅为Type2截止遗漏增加事件，两种模式分开训练评估，结果不可混表。")
p("最小测试中，当前100秒、截止3700秒、资源足够、下个完成事件6000秒，作者路径推进至6000秒，修正路径停于3700秒。这是工程与语义验证，不能计为创新。28项测试通过；CUDA短跑采用1轮、2条256作业轨迹训练，再以1条64作业序列评估两种模式，四步均exit 0。记录见 research/smoke_summary.json，仅证明链条可执行，尚未完整复现。")

page()
h("3 创新方向候选")
p("以下方向均是可执行的候选问题，尚未通过完整近邻文献核验，不声称新颖性。先限定一个问题和一条证据链，避免同时更换状态、网络和奖励后无法解释提升来源。")
h("方向一 预测误差下的稳健调度", 2)
p("手法为加误差工况，必要时再改决策接口。先只扰动策略可见的24小时绿色预测，真实供给与能源积分保持不变，测试有偏误差、随机误差和时间错位；随后考虑在绿色编码器加入不确定性输入，或调整延迟与回填的安全余量。")
p("可检验的问题是：对预测误差敏感的环节在作业选择、主动延迟还是回填接受门槛。与近邻工作的差异应落实为信息可得性、误差模型、约束和评价协议，而不是只更换网络名称。")
p("退化检查：误差幅度设为0，保持相同供给、作业段、训练种子与模式，应回到本项目已验证基线；新增保护参数退回关闭值后应恢复原决策。成本较低至中等，先做冻结策略的压力测试，再决定是否支付重新训练成本。")
h("方向二 等待尾部与服务约束", 2)
p("手法为加服务指标并改奖励或约束。平均slowdown可能掩盖少数作业的长等待，可新增等待时间或slowdown的P95、P99、最大值及违反服务阈值的比例，再评估约束PPO或拉格朗日惩罚。应先固定一项明确约束，避免把多个新指标简单加权。")
p("退化检查：约束关闭或乘子为0时，回到相同模式的原奖励与策略训练；约束阈值充分宽松时应不改变可行解。成本中等，需要更多独立种子和足够长的评估段来稳定尾部估计，也可能降低绿色比例，因此必须报告权衡曲线。")
h("方向三 跨负载泛化与迁移", 2)
p("手法为换训练协议，必要时改归一化或编码器。使用训练负载与未见测试负载的明确划分，比较单负载训练、混合训练、零样本迁移和少量微调。优先控制队列特征缩放、集群规模和供给强度，防止把分布差异误判为架构收益。")
p("退化检查：训练与测试恢复同负载、同集群和同预处理时，应回到本项目同分布基线；关闭域随机化或适配模块后恢复原训练协议。成本最高，比较矩阵和训练组合增长较快；若本地算力或数据划分不足，应排在前两项之后。")

page()
h("4 推荐方向与推进条件")
p("优先推荐方向一的第一阶段，即只做预测误差压力测试。它直接检验论文依赖未来24小时预测的假设，能够利用同一冻结策略形成配对对照，也能定位主动延迟与绿色回填何时失效。只有观察到稳定、可重复的缺口，且文献核验支持独立问题后，再确定最小算法改造。")
h("实验推进顺序", 2)
p("第一步，完成作者模式的固定负载基线，记录代码版本、配置、数据范围、种子、训练预算和权重。先以PPO及GAS-MARL为主，逐步补充FCFS、F2和GA；LPTPN按论文仅在无回填条件比较。")
p("第二步，对author与paper模式做配对验证，所有其他条件一致。把依赖、路径、参数传递等工程修复与Type2语义修正分开记录；若修正导致结果改变，先解释改变，再进行创新实验。")
p("第三步，设置彼此分离的训练、验证和测试作业段，并分别控制训练种子、评估动作采样和预测扰动种子。保留逐序列结果；随机训练以多种子均值、离散程度和配对差异为依据，不把单次落入固定1%或5%误差视为唯一成功标准。")
p("第四步，先跑无误差基线及少量误差水平，再扩展压力测试。算力与耗时先用小规模运行实测后估算；300轮完整训练及三负载多种子组合尚不能依据仓库直接承诺时长。")
h("形成研究结论前的判据", 2)
p("基线完成的条件是输入、评价口径与锚点可追溯，实际训练和评估日志完整，差异有说明。候选方向成立的条件是问题在多个种子和独立作业段上重复出现，并且与近期方法的差异明确。创新有效的条件是公平预算下改善目标或约束，退化检查通过，同时报告损失、失败工况和额外开销。")
h("主要来源与邻近工作", 2)
p("[1] Chen R, Lin W, Huang H, Ye X, Peng Z. GAS-MARL: Green-Aware job Scheduling algorithm for HPC clusters based on Multi-Action Deep Reinforcement Learning. Future Generation Computer Systems 167 (2025) 107760. DOI 10.1016/j.future.2025.107760。")
p("[2] 作者源码与数据  https://github.com/sioncr/green-rl-sched  原始审查提交 0d89df599ffe092bdd8bfb0c86a0628e58833bde。")
p("[3] SSRN 2026  Job Energy Prediction-Driven Green Scheduling for Heterogeneous HPC Workloads via Attention-Enhanced Dual-Critic Reinforcement Learning  https://papers.ssrn.com/sol3/papers.cfm?abstract_id=6530554  待全文核验的邻近工作，尚不能据标题确认方法差异。")
p("[4] Applied Energy 2026  Energy-aware job scheduling for green data centers under uncertainty: A structure-compressed spectral risk control approach  DOI 10.1016/j.apenergy.2026.128137  待全文核验的邻近工作，尚不能排除与稳健性或风险约束方向重合。")

doc.core_properties.title = "GAS MARL 模型拆解与研究方向报告"
doc.core_properties.subject = "GAS-MARL 论文拆解与候选研究方向"
doc.core_properties.author = ""
doc.save(OUT / "拆解报告_GAS_MARL.docx")
(OUT / "拆解报告_GAS_MARL.md").write_text("\n".join(markdown), encoding="utf-8")
print(OUT / "拆解报告_GAS_MARL.docx")
