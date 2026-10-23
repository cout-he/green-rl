"""Build the three-seed report from audited evidence using bundled python-docx."""
from pathlib import Path
import hashlib
import json
import statistics
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports'
EVIDENCE = ROOT / 'research/evidence/multiseed'
summary = json.loads((EVIDENCE / 'multiseed_summary.json').read_text(encoding='utf-8'))
audit = json.loads((EVIDENCE / 'training_verification.json').read_text(encoding='utf-8'))
assert audit['status'] == 'verified'
assert len(summary['per_seed']) == 24 and len(summary['aggregate']) == 8


def value(seed, algorithm, mode='author'):
    return next(r for r in summary['per_seed'] if r['seed'] == seed and r['algorithm'] == algorithm and r['metric_mode'] == mode)


def aggregate(algorithm, mode='author'):
    return next(r for r in summary['aggregate'] if r['algorithm'] == algorithm and r['metric_mode'] == mode)


marl, ppo = aggregate('MARL'), aggregate('PPO')
assert all(value(s, 'MARL')['bsld_mean'] < value(s, 'PPO')['bsld_mean'] and value(s, 'MARL')['reutil_mean'] > value(s, 'PPO')['reutil_mean'] for s in (0, 1, 2))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.6)
section.left_margin, section.right_margin = Cm(2), Cm(2)
for name, size in [('Normal', 11), ('Title', 21), ('Subtitle', 10), ('Heading 1', 14), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    style.paragraph_format.line_spacing = 1.18
    style.paragraph_format.space_after = Pt(7)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('GAS-MARL 三种子基线验证  ·  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
md = ['# GAS-MARL 三种子基线验证报告', '']


def p(text):
    result = doc.add_paragraph(text)
    md.extend([text, ''])
    return result


def heading(text, page=False):
    result = doc.add_heading(text, 1)
    result.paragraph_format.page_break_before = page
    md.extend(['## ' + text, ''])


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for values in rows:
        for cell, text in zip(t.add_row().cells, values):
            cell.text = str(text)
    for ri, row in enumerate(t.rows):
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement('w:' + edge)
                for key, value_ in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    item.set(qn('w:' + key), value_)
                borders.append(item)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge, amount in [('top', '80'), ('bottom', '80'), ('left', '90'), ('right', '90')]:
                item = OxmlElement('w:' + edge)
                item.set(qn('w:w'), amount)
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            pr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DCE6F1' if ri == 0 else ('F7F9FB' if ri % 2 == 0 else 'FFFFFF'))
            pr.append(shade)
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.12
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    run.bold = ri == 0
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    def md_cell(text):
        return str(text).replace('|', '&#124;').replace('\n', '<br>')
    md.extend(['| ' + ' | '.join(map(md_cell, headers)) + ' |', '|' + '---|' * len(headers)])
    md.extend('| ' + ' | '.join(map(md_cell, row)) + ' |' for row in rows)
    md.append('')


def figure(stem, caption):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(OUT / 'figures' / (stem + '.png')), width=Cm(16.7))
    paragraph.paragraph_format.keep_with_next = True
    result = p(caption)
    for run in result.runs:
        run.font.size = Pt(9)
    md.extend([f'![{caption}](figures/{stem}.png)', ''])


doc.add_paragraph('GAS-MARL 三种子基线验证报告', 'Title')
doc.add_paragraph('Lublin 256  Green Backfilling  Python PyTorch  2026年10月23日', 'Subtitle')
p('种子0、1、2的GAS-MARL和PPO均已完成300轮训练及统一配对评测，最后一批于北京时间10月23日16时49分结束。本报告核对六份完整训练状态并汇总可比的三种子结果，保留此前单种子报告和全部原始输出。')
heading('一 复现到什么程度')
p('在本次限定配置下，三个种子的作者模式结果均支持GAS-MARL相对PPO降低平均BSLD并提高绿电利用率。跨种子均值仍与论文不同，且只覆盖一个工作负载和一种回填策略；据此可确认本阶段基线验证完成，不能宣称整篇论文或所有表格已复现。')
rows = []
for seed in (0, 1, 2):
    a, b = value(seed, 'MARL'), value(seed, 'PPO')
    rows.append([seed, f"{a['bsld_mean']:.4f}", f"{b['bsld_mean']:.4f}", f"{100*a['reutil_mean']:.2f}", f"{100*b['reutil_mean']:.2f}"])
rows.append(['均值 ± SD', f"{marl['bsld_mean']:.4f} ± {marl['bsld_seed_sd']:.4f}", f"{ppo['bsld_mean']:.4f} ± {ppo['bsld_seed_sd']:.4f}", f"{100*marl['reutil_mean']:.2f} ± {100*marl['reutil_seed_sd']:.2f}", f"{100*ppo['reutil_mean']:.2f} ± {100*ppo['reutil_seed_sd']:.2f}"])
table(['种子', 'GAS-MARL\nBSLD ↓', 'PPO\nBSLD ↓', 'GAS-MARL\n绿电 % ↑', 'PPO\n绿电 % ↑'], rows, [2.4, 3.7, 3.7, 3.5, 3.5])
p(f"以三种子均值之比计算，GAS-MARL的BSLD比PPO低{100*(1-marl['bsld_mean']/ppo['bsld_mean']):.2f}%，ReUtil高{100*(marl['reutil_mean']-ppo['reutil_mean']):.2f}个百分点。这是复现方法与对照算法的比较，不是新算法收益。")
p('每个种子的评测点是十段1024作业结果的算术均值；表中SD是三个种子均值之间的样本标准差。训练种子与基础评测种子同号，因此这里描述端到端随机性，不能分解为纯训练方差。')
p('BSLD越低越好；ReUtil越高越好，其分母为集群总耗能。十个固定窗口有11对重叠，不能把30个窗口当作30次独立训练。本报告不进行显著性检验。')

heading('二 作者模式对照与配对差异', page=True)
figure('gas_marl_multiseed_evaluation', '图1 三种子作者模式统一配对评测。上排为绝对指标，下排为同种子GAS-MARL相对PPO的差异；每个种子点为十窗均值，菱形和误差条为三种子均值及样本SD。空心方块是论文Table 3和4锚点。120条作者模式记录全部参与统计。')
rows = []
for a in ('MARL', 'PPO'):
    r = aggregate(a)
    anchor = (31.738, .8154) if a == 'MARL' else (56.358, .6767)
    rows.append(['GAS-MARL' if a == 'MARL' else a, f'{anchor[0]:.3f}', f'{100*anchor[1]:.2f}', f"{100*(r['bsld_mean']/anchor[0]-1):+.2f}%", f"{100*(r['reutil_mean']-anchor[1]):+.2f}"])
table(['算法', '论文BSLD', '论文绿电 %', 'BSLD偏差', '绿电偏差 pp'], rows, [3, 3, 3.2, 3.5, 3.5])
p('上述偏差描述三种子均值与发表均值的差别。性能排序重复出现，并不意味着软件、采样和实现语义的偏离已被消除。')

heading('三 训练轨迹与证据核验', page=True)
figure('gas_marl_multiseed_training', '图2 六次训练的1800行原始记录全部保留。每种算法粗线为同一轮三个种子的均值，色带为样本SD，浅线为各次训练；未选择最好轮次，也未做移动平均。')
table(['核验项目', '核验范围'], [
    ['完整训练', '六份CSV均300行；日志epoch连续1至300，并与CSV逐行一致'],
    ['最终模型', '六份checkpoint均为第300轮；与独立actor和critic权重逐tensor相等，参数有限'],
    ['完整状态', '双Adam、全局及环境RNG与CSV历史齐全，元数据匹配seed和配置'],
    ['运行与来源', 'seed0原批次8步与seed1/2批次16步均正常退出；源码和输入数据SHA256核验'],
    ['统一评测', '3种子×2模式×4算法×10段，共240条记录；均值、CSV与奖励公式核验'],
    ['配对零效应', '三个种子的FCFS、F2和PPO在author与paper模式逐行一致'],
], [3.4, 13.4])
p('训练源提交分别为seed0的0b6541e和seed1/2的8915ae3；作者参考版本为0d89df5。后续工程修复和评测播种变更有单独记录，核验以实际运行保存的源码快照和输入清单为准。')
p('环境为Python 3.11.15、PyTorch 2.7.1+cu128、NumPy 1.24.4、RTX 5060。作者PyTorch 2.1.1因显卡兼容升级；NumPy重复约束统一为1.24.4，OpenCV固定4.11.0.86。这些是复现版本偏离，未计作算法创新。')

heading('四 截止语义敏感性与研究边界', page=True)
p('所有模型均按作者语义训练。paper模式只改变被测试发现的Type 2延迟截止事件处理，用同一第300轮模型和配对随机流评测；没有训练新方法。早期seed0 legacy评测存在共享随机流混杂，本报告统一使用另存的paired评测，未混入其较好数字。')
rows = []
for seed in (0, 1, 2):
    a, b = value(seed, 'MARL'), value(seed, 'MARL', 'paper')
    rows.append([seed, f"{a['bsld_mean']:.4f} → {b['bsld_mean']:.4f}", f"{100*(b['bsld_mean']/a['bsld_mean']-1):+.2f}%", f"{100*a['reutil_mean']:.2f} → {100*b['reutil_mean']:.2f}"])
table(['种子', 'BSLD  author → paper', 'BSLD变化', '绿电 %  author → paper'], rows, [1.7, 5.7, 3, 6.4])
paper = aggregate('MARL', 'paper')
p(f"paper模式GAS-MARL的跨种子BSLD为{paper['bsld_mean']:.4f} ± {paper['bsld_seed_sd']:.4f}，ReUtil为{100*paper['reutil_mean']:.2f}% ± {100*paper['reutil_seed_sd']:.2f}个百分点。种子1的BSLD恶化且绿电利用率下降，说明截止语义修正没有带来一致提升，不能把工程修正包装成算法贡献。")
p('已完成范围为Lublin-256、Green-Backfilling以及FCFS、F2、PPO和GAS-MARL的比较。论文中的GA、LPTPN、Cirne、Jann及其他回填配置尚未完成；相同轨迹上的十个窗口也不是严格隔离的外部分布测试。')
p('下一阶段先补核直接近邻，再筛选可证伪的候选。绿电预测误差下延迟与回填稳健性仍待查新，不能因为加入预测、注意力或双critic便认定创新。若开展小试，应复用固定权重、冻结真实供给与作业，仅扰动调度器可见的未来预测，并先验证零扰动完全退化。')
p('本阶段未追加长期训练。代码、图表和小型证据同步研究分支及草稿PR；权重、虚拟环境与用户论文PDF保留本地。首轮报告继续保存，自动跟进保持低频，仅在完成、失败或有实质科研发现时通知。')
p('复核入口为scripts/verify_multiseed.py与scripts/analyze_multiseed.py。研究证据目录multiseed保存逐序列数据、每种子与跨种子汇总、源输入哈希和模型核验记录，图表附源数据与绘图脚本。')

OUT.mkdir(exist_ok=True)
doc.save(OUT / '三种子基线验证报告_GAS_MARL.docx')
(OUT / '三种子基线验证报告_GAS_MARL.md').write_text('\n'.join(md), encoding='utf-8')
record = {'status': 'report_created_pending_visual_qa', 'report': 'reports/三种子基线验证报告_GAS_MARL.docx',
          'inputs_sha256': {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in (EVIDENCE / 'multiseed_summary.json', EVIDENCE / 'training_verification.json')}}
(EVIDENCE / 'report_build.json').write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(record['report'])
